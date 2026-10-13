"""Read member roster files (CSV or XLSX) into plain text rows.

Rules this module enforces:
- Size, row, and column limits are checked while reading, and every failure raises
  ImportFileError with a code the API can return in the shared error format.
- Spreadsheet formulas are never evaluated. A formula cell is kept as its literal text.
- Every value comes back as text. Cells stored as numbers are recorded in numeric_cells so
  validation can flag identifiers that may have lost their leading zeros.
"""

import csv
import io
import os
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime

# Proposed limits. They are the first thing to change if the team picks different numbers.
MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_DATA_ROWS = 2000
MAX_COLUMNS = 50
MAX_XLSX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024

SUPPORTED_EXTENSIONS = (".csv", ".xlsx")
_FORMULA_STARTS = ("=", "+", "-", "@")
_XLSX_MAGIC = b"PK\x03\x04"


class ImportFileError(Exception):
    """The file cannot be used. `code` and `message` match the shared API error shape."""

    def __init__(self, code: str, message: str, field: str | None = None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.field = field

    def to_dict(self) -> dict:
        return {"error": {"code": self.code, "message": self.message, "field": self.field}}


@dataclass(frozen=True)
class _Cell:
    text: str
    kind: str = "text"  # "text", "number", or "formula"


_EMPTY = _Cell("")


@dataclass
class ParsedRow:
    number: int  # the row's position in the file, counting the header row, as a spreadsheet shows it
    values: dict[str, str]  # column header -> cell text, "" when empty


@dataclass
class ParsedFile:
    file_type: str  # "csv" or "xlsx"
    headers: list[str]
    rows: list[ParsedRow]
    numeric_cells: set[tuple[int, str]] = field(default_factory=set)  # (row number, header)
    formula_cells: set[tuple[int, str]] = field(default_factory=set)


def parse_member_file(filename: str, content: bytes) -> ParsedFile:
    """Read an uploaded roster file. Raises ImportFileError for any problem with the file."""
    if len(content) == 0:
        raise ImportFileError("EMPTY_FILE", "The file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(
            "FILE_TOO_LARGE",
            f"The file is larger than the {MAX_FILE_BYTES // (1024 * 1024)} MB limit.",
        )

    extension = os.path.splitext(filename or "")[1].lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ImportFileError("UNSUPPORTED_FILE_TYPE", "Upload a .csv or .xlsx file.")

    if extension == ".csv":
        if content.startswith(_XLSX_MAGIC):
            raise ImportFileError(
                "UNREADABLE_FILE", "This is a spreadsheet saved with a .csv name. Upload it as .xlsx."
            )
        return _assemble("csv", _csv_rows(content))

    if not content.startswith(_XLSX_MAGIC):
        raise ImportFileError("UNREADABLE_FILE", "This file is not a valid .xlsx spreadsheet.")
    return _assemble("xlsx", _xlsx_rows(content))


# ---------------------------------------------------------------- shared assembly


def _assemble(file_type: str, raw_rows: Iterable[tuple[int, list[_Cell]]]) -> ParsedFile:
    headers: list[str] | None = None
    rows: list[ParsedRow] = []
    numeric_cells: set[tuple[int, str]] = set()
    formula_cells: set[tuple[int, str]] = set()
    scanned = 0
    scan_limit = MAX_DATA_ROWS * 2 + 10  # bounds files padded with thousands of blank rows

    for number, cells in raw_rows:
        scanned += 1
        if scanned > scan_limit:
            raise ImportFileError("TOO_MANY_ROWS", f"The file has more than {MAX_DATA_ROWS} member rows.")
        if all(not cell.text.strip() for cell in cells):
            continue

        if headers is None:
            headers = _check_headers([cell.text.strip() for cell in cells])
            continue

        if len(rows) >= MAX_DATA_ROWS:
            raise ImportFileError("TOO_MANY_ROWS", f"The file has more than {MAX_DATA_ROWS} member rows.")

        values: dict[str, str] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            cell = cells[index] if index < len(cells) else _EMPTY
            values[header] = cell.text
            if cell.kind == "number":
                numeric_cells.add((number, header))
            elif cell.kind == "formula":
                formula_cells.add((number, header))
        rows.append(ParsedRow(number=number, values=values))

    if headers is None:
        raise ImportFileError("MISSING_HEADER", "The file has no header row.")
    if not rows:
        raise ImportFileError("EMPTY_FILE", "The file has a header row but no member rows.")
    return ParsedFile(file_type, [h for h in headers if h], rows, numeric_cells, formula_cells)


def _check_headers(headers: list[str]) -> list[str]:
    named = [h for h in headers if h]
    if not named:
        raise ImportFileError("MISSING_HEADER", "The first row must contain column names.")
    if len(named) > MAX_COLUMNS:
        raise ImportFileError("TOO_MANY_COLUMNS", f"The file has more than {MAX_COLUMNS} columns.")
    seen: set[str] = set()
    for header in named:
        key = header.casefold()
        if key in seen:
            raise ImportFileError(
                "DUPLICATE_COLUMN", f'The column name "{header}" appears more than once.', field=header
            )
        seen.add(key)
    return headers


# ---------------------------------------------------------------- CSV


def _decode_csv(content: bytes) -> str:
    if b"\x00" in content:
        raise ImportFileError("UNREADABLE_FILE", "This file is not a readable CSV.")
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        return content.decode("cp1252")  # what Excel on Windows often writes
    except UnicodeDecodeError:
        raise ImportFileError("UNREADABLE_FILE", "The file's text encoding is not supported.") from None


def _csv_rows(content: bytes) -> Iterator[tuple[int, list[_Cell]]]:
    text = _decode_csv(content)
    first_line = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = max(",;\t", key=first_line.count) if any(d in first_line for d in ",;\t") else ","
    try:
        for number, record in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter), start=1):
            yield number, [_Cell(value, "formula" if value.lstrip().startswith(_FORMULA_STARTS) else "text") for value in record]
    except csv.Error:
        raise ImportFileError("UNREADABLE_FILE", "The CSV file could not be read.") from None


# ---------------------------------------------------------------- XLSX


def _xlsx_rows(content: bytes) -> Iterator[tuple[int, list[_Cell]]]:
    _check_zip(content)
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception:
        raise ImportFileError("UNREADABLE_FILE", "This file is not a valid .xlsx spreadsheet.") from None

    try:
        sheet = workbook.worksheets[0]  # only the first sheet is read
        for number, row in enumerate(sheet.iter_rows(), start=1):
            yield number, [_xlsx_cell(cell) for cell in row]
    except ImportFileError:
        raise
    except Exception:
        raise ImportFileError("UNREADABLE_FILE", "The spreadsheet could not be read.") from None
    finally:
        workbook.close()


def _check_zip(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        raise ImportFileError("UNREADABLE_FILE", "This file is not a valid .xlsx spreadsheet.") from None
    if "xl/workbook.xml" not in names:
        raise ImportFileError("UNREADABLE_FILE", "This file is not a valid .xlsx spreadsheet.")
    if total > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ImportFileError("FILE_TOO_LARGE", "The spreadsheet is too large once opened.")


def _xlsx_cell(cell) -> _Cell:
    value = getattr(cell, "value", None)
    if value is None:
        return _EMPTY
    if getattr(cell, "data_type", None) == "f":
        return _Cell(str(value), "formula")  # kept as text, never calculated
    if isinstance(value, datetime):
        return _Cell(value.date().isoformat())
    if isinstance(value, date):
        return _Cell(value.isoformat())
    if isinstance(value, bool):
        return _Cell(str(value))
    if isinstance(value, int):
        return _Cell(str(value), "number")
    if isinstance(value, float):
        text = str(int(value)) if value.is_integer() else repr(value)
        return _Cell(text, "number")
    return _Cell(str(value))
